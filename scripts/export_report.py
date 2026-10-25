"""Build a private report from the retained ICTAK template with bundled python-docx."""
from copy import deepcopy
from io import BytesIO
import hashlib
import json
from pathlib import Path
import re
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parents[1]
REFERENCE = ROOT / 'Summer Internship Report __ React.docx'
OUTPUT = ROOT / '.local/submission/LabFlow_Project_Report.docx'
DETAILS = ROOT / '.local/submission-details.json'
details = json.loads(DETAILS.read_text()) if DETAILS.exists() else {}
assert hashlib.sha256(REFERENCE.read_bytes()).hexdigest() == '192f563bee16a0853fffd4924524d666cdc678b6d8aa346fc65eb22a51f5fb90'
doc = Document(REFERENCE)
source = [deepcopy(p._p) for p in doc.paragraphs]


def replace_text(paragraph, text):
    first = paragraph.runs[0] if paragraph.runs else paragraph.add_run()
    properties = deepcopy(first._r.rPr) if first._r.rPr is not None else None
    for child in list(paragraph._p):
        if child.tag != qn('w:pPr'):
            paragraph._p.remove(child)
    run = paragraph.add_run(text)
    if properties is not None:
        run._r.insert(0, properties)
    return paragraph


for index, text in {
    7: 'LabFlow',
    8: 'Laboratory Equipment Management',
    13: f"{details.get('fullName', '[Full name]')}\nRegister No: {details.get('studentId', '[Student ID]')} | Sr No: {details.get('serialNumber', '[Serial number]')}",
    14: details.get('institution', '[Institution name]'),
    15: 'Prepared 25 October 2026',
}.items():
    replace_text(doc.paragraphs[index], text)

body = doc.element.body
for child in list(body)[17:]:
    if child.tag != qn('w:sectPr'):
        body.remove(child)
# The cover's explicit break paragraph can itself overflow after individual metadata.
# The first body heading already starts a new page, so retain only that one break.
body.remove(doc.paragraphs[16]._p)


def paragraph(text, kind='body'):
    node = deepcopy(source[17 if kind == 'heading' else 23 if kind == 'numbered' else 18])
    body.insert(len(body) - 1, node)
    p = replace_text(Paragraph(node, doc._body), text)
    p.paragraph_format.widow_control = True
    p.paragraph_format.keep_with_next = kind == 'heading'
    return p


def clean(text):
    text = re.sub(r'\[([^\]]+)\]\(([^)]+)\)', r'\1 (\2)', text)
    return text.replace('**', '').replace('`', '')


def figure(path, caption):
    from PIL import Image
    with Image.open(path) as image:
        width, height = image.size
    p = paragraph(caption)
    p.paragraph_format.page_break_before = True
    p.paragraph_format.keep_with_next = True
    size = min(6.5, 7.5 * width / height)
    holder = paragraph('')
    holder.paragraph_format.keep_together = True
    holder.add_run().add_picture(str(path), width=Inches(size))


def table(lines):
    rows = [[clean(cell.strip()) for cell in line.strip('|').split('|')] for line in lines]
    rows = [row for row in rows if not all(re.fullmatch(r':?-+:?', cell) for cell in row)]
    widths = [1.35, 2.95, 2.2]
    tbl = doc.add_table(rows=0, cols=len(rows[0]))
    tbl.autofit = False
    for col, width in zip(tbl.columns, widths):
        col.width = Inches(width)
    borders = OxmlElement('w:tblBorders')
    for side in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        el = OxmlElement('w:' + side)
        for key, value in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
            el.set(qn('w:' + key), value)
        borders.append(el)
    tbl._tbl.tblPr.append(borders)
    for index, row in enumerate(rows):
        cells = tbl.add_row().cells
        for cell, value, width in zip(cells, row, widths):
            cell.width = Inches(width)
            cell.vertical_alignment = 1
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(5)
            p.paragraph_format.space_before = Pt(5)
            for run in [p.add_run(value)]:
                run.font.name = 'Roboto Condensed'
                run.font.size = Pt(10.5)
                run.bold = index == 0
            if index == 0:
                shading = OxmlElement('w:shd')
                shading.set(qn('w:fill'), 'E7F0F8')
                cell._tc.get_or_add_tcPr().append(shading)
        props = tbl.rows[-1]._tr.get_or_add_trPr()
        props.append(OxmlElement('w:cantSplit'))
        if index == 0:
            props.append(OxmlElement('w:tblHeader'))


paragraph('EXECUTIVE SUMMARY', 'heading').paragraph_format.page_break_before = True
lines = (ROOT / 'docs/REPORT_DRAFT.md').read_text().splitlines()
index = 0
active = False
section = ''
figure_count = 0
while index < len(lines):
    line = lines[index].strip()
    if line.startswith('## '):
        section = line[3:]
        if active:
            paragraph('RESULTS & FINDINGS' if section == 'Results and findings' else section.upper(), 'heading')
        active = True
    elif active and line:
        match = re.fullmatch(r'!\[([^]]+)\]\(([^)]+)\)', line)
        if match:
            figure_count += 1
            figure(ROOT / 'docs' / match[2], f'Figure {figure_count}. {match[1]}.')
        elif line.startswith('|'):
            block = []
            while index < len(lines) and lines[index].strip().startswith('|'):
                block.append(lines[index].strip())
                index += 1
            table(block)
            continue
        elif line.startswith('- '):
            paragraph(clean(line[2:]), 'numbered' if section == 'Objectives' else 'body')
        else:
            paragraph(clean(line))
    index += 1

screenshots = ROOT / '.local/submission/screenshots'
for filename, caption in [
    ('01-login.jpg', 'Sign-in screen and laboratory access information.'),
    ('03-student-catalog.jpg', 'Student equipment catalog with availability and search filters.'),
    ('04-student-request-other.jpg', 'Equipment and purpose dropdowns with Other descriptions.'),
    ('05-student-request-history.jpg', 'Student borrowing history and request states.'),
    ('06-admin-dashboard.jpg', 'Administrator circulation summary and inventory state counts.'),
    ('08-admin-register-asset.jpg', 'Administrator asset registration with preset metadata choices.'),
    ('09-admin-request-management.jpg', 'Administrator request management with status, overdue and damage filters.'),
    ('10-admin-return-and-history.jpg', 'Inspected return, original requested description and retained lifecycle timeline.'),
]:
    figure_count += 1
    figure(screenshots / filename, f'Figure {figure_count}. {caption} LabFlow application interface, 25 October 2026.')

p = paragraph(f'Figure {figure_count + 1}. Mobile request form: upper view and scrolled lower fields. LabFlow application interface, 25 October 2026.')
p.paragraph_format.page_break_before = True
p.paragraph_format.keep_with_next = True
phones = doc.add_table(rows=1, cols=2)
phones.autofit = False
for col in phones.columns:
    col.width = Inches(3.25)
for cell, filename in zip(phones.rows[0].cells, ['11-mobile-overview.jpg', '12-mobile-request-form.jpg']):
    cell.width = Inches(3.25)
    cell.paragraphs[0].add_run().add_picture(str(screenshots / filename), width=Inches(3.0))

buffer = BytesIO()
doc.save(buffer)
modified = {'word/document.xml', 'word/_rels/document.xml.rels', '[Content_Types].xml'}
OUTPUT.parent.mkdir(parents=True, exist_ok=True)
with ZipFile(REFERENCE) as original, ZipFile(buffer) as built, ZipFile(OUTPUT, 'w', ZIP_DEFLATED) as final:
    for name in original.namelist():
        final.writestr(name, built.read(name) if name in modified else original.read(name))
    for name in built.namelist():
        if name not in original.namelist() and name.startswith('word/media/'):
            final.writestr(name, built.read(name))
with ZipFile(REFERENCE) as original, ZipFile(OUTPUT) as result:
    assert all(original.read(name) == result.read(name) for name in original.namelist() if name not in modified)
print(OUTPUT)
