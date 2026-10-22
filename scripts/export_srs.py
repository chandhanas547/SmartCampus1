"""Export the Markdown SRS with python-docx; run using the bundled document runtime."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'docs/SRS_DRAFT.md'
OUTPUT = ROOT / 'docs/exports/LabFlow_SRS_DRAFT_0.4.docx'
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = Inches(.7)
section.left_margin = section.right_margin = Inches(.8)
section.header_distance = section.footer_distance = Inches(.3)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3']:
    style = doc.styles[name]
    style.font.name = 'Arial'
    style.font.color.rgb = RGBColor(0, 0, 0)
for border in list(doc.styles.element.iter(qn('w:pBdr'))):
    border.getparent().remove(border)
doc.styles['Normal'].font.size = Pt(11)
doc.styles['Normal'].paragraph_format.space_after = Pt(6)
doc.styles['Normal'].paragraph_format.line_spacing = 1.08
doc.styles['Title'].font.size = Pt(25)
for name, size in [('Heading 1', 16), ('Heading 2', 13), ('Heading 3', 11.5)]:
    doc.styles[name].font.size = Pt(size)
    doc.styles[name].paragraph_format.space_before = Pt(12)
    doc.styles[name].paragraph_format.space_after = Pt(6)
header = section.header.paragraphs[0]
header.text = 'LabFlow SRS 0.4 Draft'
header.runs[0].font.size = Pt(9)
footer = section.footer.paragraphs[0]
footer.alignment = 2
footer.add_run('Page ').font.size = Pt(9)
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)

def inline(paragraph, text):
    text = re.sub(r'\[([^\]]+)\]\(([^)]+)\)', r'\1 (\2)', text)
    for token in re.split(r'(\*\*.*?\*\*|`[^`]+`)', text):
        run = paragraph.add_run(token.strip('*`') if token.startswith(('**', '`')) else token)
        if token.startswith('**'): run.bold = True

def table(lines):
    rows = [[cell.strip() for cell in line.strip().strip('|').split('|')] for line in lines]
    rows = [row for row in rows if not all(re.fullmatch(r':?-+:?', cell) for cell in row)]
    count = len(rows[0])
    tbl = doc.add_table(rows=0, cols=count)
    tbl.autofit = False
    width = 6.9
    if count == 2: widths = [1.8, 5.1]
    elif count == 3: widths = [1.05, 1.55, 4.3]
    elif count == 4: widths = [.65, 1.15, 2.55, 2.55]
    elif count == 5: widths = [1.6, .7, 2.2, 1.1, 1.3]
    else: widths = [width / count] * count
    if rows[0][0] == 'Version': widths = [.7, 1, .8, 4.4]
    if count == 4 and rows[0][0] == 'Capability': widths = [2.1, .6, 2.1, 2.1]
    if count == 3 and rows[0][0] == 'Capability': widths = [1.2, 3.7, 2]
    if count == 4 and rows[0][0] == 'Action': widths = [1.15, 1.25, 2.25, 2.25]
    if count == 2 and rows[0][0] in ['Assumption', 'Decision']: widths = [4.8, 2.1]
    for col, value in zip(tbl.columns, widths): col.width = Inches(value)
    for index, row in enumerate(rows):
        cells = tbl.add_row().cells
        for cell, value, size in zip(cells, row, widths):
            cell.width = Inches(size)
            cell.vertical_alignment = 1
            props = cell._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for side in ['top', 'bottom', 'left', 'right']:
                item = OxmlElement('w:' + side); item.set(qn('w:w'), '85'); item.set(qn('w:type'), 'dxa'); margins.append(item)
            props.append(margins)
            para = cell.paragraphs[0]
            para.paragraph_format.space_after = Pt(3)
            para.paragraph_format.space_before = Pt(3)
            para.paragraph_format.line_spacing = 1.05
            inline(para, value)
            for run in para.runs:
                run.font.size = Pt(10.5)
                if index == 0: run.bold = True
            if index == 0:
                shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'EEEEEE'); props.append(shade)
        tr_props = tbl.rows[-1]._tr.get_or_add_trPr()
        tr_props.append(OxmlElement('w:cantSplit'))
        if index == 0: tr_props.append(OxmlElement('w:tblHeader'))
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

lines = SOURCE.read_text().splitlines()
index = 0
while index < len(lines):
    line = lines[index]
    if not line.strip(): index += 1; continue
    if line.startswith('```'):
        block = []
        index += 1
        while index < len(lines) and not lines[index].startswith('```'):
            block.append(lines[index]); index += 1
        para = doc.add_paragraph()
        para.paragraph_format.keep_together = True
        run = para.add_run('\n'.join(block))
        run.font.name = 'Courier New'; run.font.size = Pt(10)
        index += 1
        continue
    if line.startswith('|'):
        block = []
        while index < len(lines) and lines[index].startswith('|'):
            block.append(lines[index]); index += 1
        table(block); continue
    if line.startswith('# '): doc.add_paragraph('LabFlow Software Requirements Specification', 'Title')
    elif re.match(r'^#{2,4} ', line):
        marker, heading = line.split(' ', 1)
        heading = re.sub(r'[^\w\s]', ' ', heading)
        doc.add_heading(re.sub(r'\s+', ' ', heading).strip(), len(marker) - 1)
    elif line.startswith('- '): inline(doc.add_paragraph(style='List Bullet'), line[2:])
    elif re.match(r'^\d+\. ', line): inline(doc.add_paragraph(), line)
    else: inline(doc.add_paragraph(), line)
    index += 1
doc.core_properties.title = 'LabFlow Software Requirements Specification'
doc.core_properties.subject = 'Revision 0.4 Draft pending approval and submission'
doc.core_properties.author = ''
for border in list(doc.element.iter(qn('w:pBdr'))):
    border.getparent().remove(border)
OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
